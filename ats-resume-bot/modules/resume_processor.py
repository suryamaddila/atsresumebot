import os
from docx import Document
from fpdf import FPDF
import PyPDF2

async def analyze_and_generate_pdf(resume_path, job_description, output_pdf_path):
    ext = resume_path.split('.')[-1]
    if ext == 'pdf':
        text = extract_text_from_pdf(resume_path)
    elif ext in ['docx', 'doc']:
        text = extract_text_from_docx(resume_path)
    else:
        text = "Unable to parse resume"
    matched_score = get_matching_score(text, job_description)
    generate_pdf(text, job_description, matched_score, output_pdf_path)

def extract_text_from_pdf(file_path):
    with open(file_path, 'rb') as file:
        reader = PyPDF2.PdfReader(file)
        text = ""
        for page in reader.pages:
            text += page.extract_text()
        return text

def extract_text_from_docx(file_path):
    doc = Document(file_path)
    return "\n".join([para.text for para in doc.paragraphs])

def get_matching_score(resume_text, job_desc):
    resume_words = set(resume_text.lower().split())
    job_words = set(job_desc.lower().split())
    matches = resume_words.intersection(job_words)
    score = len(matches) / len(job_words) if job_words else 0
    return round(score * 100, 2)

def generate_pdf(resume_text, job_desc, score, output_path):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=12)
    pdf.cell(200, 10, txt="Optimized Resume", ln=True, align='C')
    pdf.ln(10)
    pdf.multi_cell(0, 10, f"🎯 ATS Match Score: {score}%\n\nResume:\n{resume_text[:3000]}")
    pdf.output(output_path)
